from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('AI_Chat_Technical_Flow_Document.docx')

NAVY = '1F3B5B'
PALE = 'EDF3F8'
GRID = 'D9D9D9'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '6')
        tag.set(qn('w:color'), color)
        borders.append(tag)
    tc_pr.append(borders)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_font(run, size=10.5, bold=False, color='000000', name='Aptos'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text='', style=None, before=0, after=6, bold_lead=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text), 14 if level == 1 else 11.5, bold=True)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(text))
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, value in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ''
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(value), 9.5, bold=True, color='FFFFFF')
        if widths: cell.width = Inches(widths[i])
    for row_num, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ''
            if row_num % 2 == 1:
                set_cell_shading(cells[i], PALE)
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), 9.2)
            if widths: cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def code_block(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    set_cell_shading(cell, 'F3F5F7')
    set_cell_border(cell, 'C8D0D8')
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(text.split('\n')):
        if index:
            p.add_run().add_break()
        set_font(p.add_run(line), 8.7, name='Consolas')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
set_font(title.add_run('AI Chat Application Technical Flow'), 22, bold=True)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(20)
set_font(subtitle.add_run('React and Vite implementation documentation'), 11, color='404040')

heading(doc, '1 Project Summary')
paragraph(doc, 'The AI Chat application is a single page React application that provides multi-conversation chat, local persistence, and live streamed AI responses. The interface is built with React 19, Vite, Tailwind CSS, and the Groq OpenAI-compatible chat completion API. The project demonstrates component composition, reducer-driven state management, asynchronous streaming, browser storage, and user-facing recovery controls.')
paragraph(doc, 'The implemented default model is openai/gpt-oss-20b. This model was selected after verifying that it is available to the configured Groq account. The model identifier can be changed in src/hooks/useChatStream.js if account access changes.', bold_lead='The implemented default model is ')

heading(doc, '2 Technology Stack')
table(doc, ['Layer', 'Technology', 'Purpose'], [
    ['User interface', 'React 19', 'Renders the component tree and updates the view from application state.'],
    ['Build and development', 'Vite 8', 'Provides the development server, module bundling, and production build.'],
    ['Styling', 'Tailwind CSS 4', 'Applies the responsive dark chat layout and component states.'],
    ['State', 'useReducer and custom hooks', 'Maintains conversations, messages, active selection, streaming status, and persistence.'],
    ['AI service', 'Groq Chat Completions API', 'Generates streamed assistant responses through an OpenAI-compatible endpoint.'],
    ['Persistence', 'Browser localStorage', 'Stores conversations and messages so they remain after refresh.'],
])

heading(doc, '3 High Level Architecture')
code_block(doc, 'Browser\n  |\n  +-- App.jsx\n       |-- Sidebar -> ConversationItem\n       |-- ChatWindow -> MessageList -> MessageBubble -> MarkdownMessage\n       |                 +-> ChatInput\n       |\n       +-- chatReducer.js            application state transitions\n       +-- usePersistedReducer       localStorage load and save\n       +-- useChatStream             Groq request and Server Sent Event parsing\n                                   |\n                                   +--> https://api.groq.com/openai/v1/chat/completions')
paragraph(doc, 'App.jsx owns the shared chat state and passes data and callback functions downward. Presentational components do not edit shared state directly. User actions move upward through callbacks to App, which dispatches reducer actions or starts a streaming request.')

heading(doc, '4 Application State Model')
paragraph(doc, 'The reducer stores a list of conversations and the identifier of the selected conversation. The active conversation is derived with find rather than being duplicated in state.')
code_block(doc, "{\n  conversations: [{\n    id, title, createdAt, updatedAt,\n    messages: [{ id, role, content, createdAt, status, error? }]\n  }],\n  activeConversationId\n}")
table(doc, ['Action', 'Trigger', 'Reducer Result'], [
    ['NEW_CONVERSATION', 'New Chat button', 'Creates an empty conversation, prepends it, and makes it active.'],
    ['SELECT_CONVERSATION', 'Conversation item click', 'Changes only activeConversationId.'],
    ['DELETE_CONVERSATION', 'Delete icon click', 'Removes the conversation and selects the first remaining one if needed.'],
    ['ADD_MESSAGE', 'User sends input', 'Appends the user message, updates the timestamp, and titles a new conversation from its first message.'],
    ['START_ASSISTANT_MESSAGE', 'Streaming starts', 'Appends an empty assistant message with streaming status.'],
    ['APPEND_TOKEN', 'Valid SSE token arrives', 'Immutably appends content to the matching assistant message.'],
    ['FINISH_ASSISTANT_MESSAGE', 'Stream completes or is stopped', 'Changes the assistant message status to complete.'],
    ['FAIL_ASSISTANT_MESSAGE', 'Network or API failure', 'Marks the assistant message as error and stores a readable error message.'],
])

heading(doc, '5 Message and Streaming Flow')
paragraph(doc, 'The following flow occurs after the user submits text. The interface shows the user message immediately and then adds a blank assistant bubble so the user receives feedback while the API request is active.')
code_block(doc, '1. ChatInput validates trimmed text and calls onSend(content).\n2. App creates a user message with id, timestamp, role user, and complete status.\n3. App dispatches ADD_MESSAGE and passes chat history plus the new user message to useChatStream.\n4. useChatStream dispatches START_ASSISTANT_MESSAGE with status streaming.\n5. fetch sends POST /chat/completions with model, history, stream true, and Authorization header.\n6. response.body.getReader reads the response incrementally.\n7. TextDecoder converts byte chunks to text; complete SSE records are split on blank lines.\n8. Each data record is parsed and delta.content is dispatched as APPEND_TOKEN.\n9. React re-renders the assistant bubble as tokens arrive.\n10. Completion dispatches FINISH_ASSISTANT_MESSAGE; errors dispatch FAIL_ASSISTANT_MESSAGE.')

heading(doc, '6 Streaming Control and Error Handling')
table(doc, ['Situation', 'Handling', 'User Experience'], [
    ['A response is generating', 'A single AbortController is stored in a ref and isStreaming is true.', 'Send is disabled and the Stop button is shown.'],
    ['User selects Stop', 'The controller aborts the fetch request.', 'The partial assistant response remains and is marked complete.'],
    ['Missing API key', 'The hook throws a clear configuration error before fetch.', 'The assistant bubble displays the error and Retry control.'],
    ['API, network, or parsing failure', 'The hook catches the error and dispatches FAIL_ASSISTANT_MESSAGE.', 'The error is displayed inline; Retry resends from the latest user message.'],
    ['User scrolls upward', 'useAutoScroll records whether the message list is near the bottom.', 'Automatic scrolling pauses until the user returns near the bottom.'],
])

heading(doc, '7 Persistence and Interface Behavior')
paragraph(doc, 'usePersistedReducer initializes the reducer from localStorage key ai-chat-state. If stored JSON is absent or invalid, it uses the initial empty state. A useEffect serializes each state update back to localStorage. Therefore conversations, titles, messages, timestamps, and the active conversation survive a browser refresh.')
bullet(doc, 'ChatInput is controlled local state. Enter sends a message; Shift Enter creates a new line.')
bullet(doc, 'MessageBubble distinguishes user and assistant roles, provides timestamps, copy feedback, and streaming/error status.')
bullet(doc, 'MarkdownMessage provides safe lightweight display for headings, bold text, inline code, and HTTPS links without rendering raw HTML.')
bullet(doc, 'Sidebar supports create, select, and delete actions. Delete stops click propagation so it does not also select the conversation.')

heading(doc, '8 Configuration and Run Instructions')
paragraph(doc, 'The project reads the API key from VITE_GROQ_API_KEY. Copy .env.example to .env.local, insert a valid Groq key, and restart the Vite development server. Vite exposes VITE-prefixed values to browser code; this is acceptable for local learning only. A production deployment should call Groq from a protected server-side route so the key never reaches the browser.', bold_lead='The project reads the API key from ')
code_block(doc, 'VITE_GROQ_API_KEY=gsk_your_groq_key')
table(doc, ['Command', 'Purpose'], [
    ['npm run dev', 'Starts the Vite development server.'],
    ['npm run lint', 'Runs ESLint across the project.'],
    ['npm run build', 'Creates and validates the production bundle.'],
])

heading(doc, '9 Current Verification')
paragraph(doc, 'The application source was verified with npm run lint and npm run build after the streaming model was changed to openai/gpt-oss-20b. The production build completed successfully, transforming 28 modules. Live responses additionally require a valid Groq key and a model enabled for that key.')

heading(doc, '10 Future Improvements')
bullet(doc, 'Move API calls to a backend or serverless function to protect API credentials.')
bullet(doc, 'Replace the lightweight markdown component with a full Markdown parser and GitHub Flavored Markdown support.')
bullet(doc, 'Add conversation search, rename, export, theme selection, and keyboard shortcut support.')
bullet(doc, 'Add automated unit tests for reducer actions and integration tests for stream parsing.')

doc.core_properties.title = 'AI Chat Application Technical Flow'
doc.core_properties.subject = 'Technical implementation and process flow documentation'
doc.core_properties.author = 'AI Chat Project Team'
doc.save(OUT)
print(OUT.resolve())
